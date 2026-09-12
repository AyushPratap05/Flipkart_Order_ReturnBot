"""
Reads the raw reference sheet (multi-item rows, links crammed into one cell)
and produces a clean tasks.xlsx with ONE ROW PER SKU/product link, ready for
the agent to process — matching the brief's "each SKU is its own row" rule.

Usage:
    python build_tasks_from_raw.py "Faym Status Test Orders.xlsx" tasks.xlsx
"""
import sys
import re
from openpyxl import Workbook, load_workbook

URL_PATTERN = re.compile(r'https?://\S+')

OUTPUT_HEADERS = [
    "Platform", "Order Id", "Product Link", "Return Window",
    "Status", "Refund ID", "Return Status", "Refund Amount", "Timestamp", "Log"
]


def extract_links(cell_value):
    """Pull every product URL out of a cell that may contain 1 or many,
    jammed together with WhatsApp timestamps/names/other junk text."""
    if not cell_value:
        return []
    text = str(cell_value)
    links = URL_PATTERN.findall(text)
    # dl.flipkart.com short-links and www.flipkart.com full links both count;
    # dedupe while preserving order
    seen = set()
    unique_links = []
    for link in links:
        link = link.rstrip('"\'),.')  # strip trailing punctuation caught by regex
        if link not in seen:
            seen.add(link)
            unique_links.append(link)
    return unique_links


def main(input_path, output_path):
    src_wb = load_workbook(input_path)
    src_ws = src_wb.active

    headers = [str(c.value).strip() if c.value else "" for c in src_ws[1]]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}

    def get(row, name):
        idx = col_idx.get(name)
        return row[idx - 1] if idx else None

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.append(OUTPUT_HEADERS)

    total_rows_in = 0
    total_skus_out = 0

    for row in src_ws.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        total_rows_in += 1

        order_id = get(row, "Order Id")
        platform = get(row, "Platform") or "Flipkart"
        return_window = get(row, "Return Window")
        product_link_cell = get(row, "Product Link")

        links = extract_links(product_link_cell)
        if not links:
            # fallback: cell might just be a plain (non-http) product name/text
            links = [str(product_link_cell)] if product_link_cell else []

        for link in links:
            out_ws.append([
                platform, order_id, link, return_window,
                "To Do", "", "", "", "", ""
            ])
            total_skus_out += 1

    out_wb.save(output_path)
    print(f"Read {total_rows_in} order rows -> wrote {total_skus_out} line-item (SKU) rows to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_tasks_from_raw.py <input.xlsx> <output.xlsx>")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
